# Genera marketing/registro-pedidos.xlsx: registro de pedidos con conteo automático de sellos por WhatsApp.
# Uso: python scripts/registro-pedidos.py
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.comments import Comment
from openpyxl.utils import get_column_letter
import json

marca = json.load(open('src/config/marca.json', encoding='utf-8'))
entregas = json.load(open('src/data/entregas.json', encoding='utf-8'))
promo = marca['promocion']

ARIAL = 'Arial'
def f(bold=False, color='000000', size=10): return Font(name=ARIAL, bold=bold, color=color, size=size)
AZUL = '0000FF'      # datos que tú capturas
AMARILLO = PatternFill('solid', fgColor='FFFF00')   # celdas para llenar
CABECERA = PatternFill('solid', fgColor='1A1433')
borde = Border(*(Side(style='thin', color='999999'),) * 4)

wb = Workbook()

# ---------- Config ----------
cfg = wb.active; cfg.title = 'Config'
cfg['A1'] = 'Configuración de la promoción'; cfg['A1'].font = f(True, size=12)
filas = [
    ('Sellos para llavero con nombre gratis', promo['sellos_llavero'], 'Número de pedido que gana el llavero con nombre.'),
    ('Sellos para abrir la cápsula', promo['sellos_capsula'], 'Número de pedido que gana la cápsula sorpresa.'),
    ('Piezas para domicilio gratis en Santa Catarina', entregas['domicilio_gratis_desde'], 'Fuente: política de entregas del sitio (src/data/entregas.json).'),
    ('Envío a otros municipios del área metropolitana ($)', entregas['envio_area_metropolitana'], 'Fuente: política de entregas del sitio.'),
]
for i, (etq, val, nota) in enumerate(filas, start=2):
    cfg.cell(i, 1, etq).font = f()
    c = cfg.cell(i, 2, val); c.font = f(color=AZUL); c.fill = AMARILLO
    cfg.cell(i, 3, nota).font = f(color='666666')
cfg['A8'] = 'Leyenda'; cfg['A8'].font = f(True)
cfg['A9'] = 'Celdas amarillas con texto azul: las llenas tú. Todo lo demás se calcula solo.'; cfg['A9'].font = f()
cfg['A10'] = 'Solo cuentan para sellos los pedidos con Tipo = Venta y Origen = Página (así lo dice la promo en el sitio).'; cfg['A10'].font = f()
cfg['A11'] = 'Cuando canjeas un premio, registra una fila con Tipo = Canje llavero o Canje cápsula: no suma sello y marca el premio como entregado.'; cfg['A11'].font = f()
cfg.column_dimensions['A'].width = 52; cfg.column_dimensions['B'].width = 10; cfg.column_dimensions['C'].width = 70

# ---------- Pedidos ----------
ped = wb.create_sheet('Pedidos')
cab = ['Fecha', 'WhatsApp', 'Nombre', 'Tipo', 'Producto', 'Piezas', 'Total ($)', 'Entrega', 'Origen', 'Sello']
for j, h in enumerate(cab, start=1):
    c = ped.cell(1, j, h); c.font = f(True, 'FFFFFF'); c.fill = CABECERA; c.alignment = Alignment(horizontal='center'); c.border = borde
ped['J1'].comment = Comment('Fórmula: 1 si Tipo = Venta y Origen = Página; si no, 0.', 'LlaveCorp')
ejemplos = [
    ('2026-10-03', '528112345678', 'Ana', 'Venta', 'Casete (con NFC)', 1, 110, 'Soriana Colosio', 'Página'),
    ('2026-10-10', '528112345678', 'Ana', 'Venta', 'Cartucho 8 bits', 2, 150, 'Cruz Roja', 'Página'),
    ('2026-10-11', '528187654321', 'Luis', 'Venta', 'Onigiri', 1, 60, 'Tianguis', 'Tianguis'),
]
N = 500
for i in range(2, N + 2):
    for j in range(1, 10):
        c = ped.cell(i, j); c.font = f(color=AZUL); c.border = borde
    ped.cell(i, 10, f'=IF(AND(D{i}="Venta",I{i}="Página"),1,0)').font = f()
    ped.cell(i, 10).border = borde
for i, fila in enumerate(ejemplos, start=2):
    for j, v in enumerate(fila, start=1):
        ped.cell(i, j, v)
ped['A2'].comment = Comment('Filas de ejemplo: bórralas cuando empieces a capturar.', 'LlaveCorp')
for col, w in zip('ABCDEFGHIJ', [12, 15, 18, 16, 26, 8, 10, 20, 11, 7]): ped.column_dimensions[col].width = w
ped.column_dimensions['G'].number_format = '$#,##0'
for i in range(2, N + 2): ped.cell(i, 7).number_format = '$#,##0'
dv_tipo = DataValidation(type='list', formula1='"Venta,Canje llavero,Canje cápsula"', allow_blank=True)
dv_ent = DataValidation(type='list', formula1='"Soriana Colosio,Cruz Roja,Domicilio,Envío,Tianguis"', allow_blank=True)
dv_org = DataValidation(type='list', formula1='"Página,Tianguis"', allow_blank=True)
for dv, col in ((dv_tipo, 'D'), (dv_ent, 'H'), (dv_org, 'I')):
    ped.add_data_validation(dv); dv.add(f'{col}2:{col}{N + 1}')
ped.freeze_panes = 'A2'

# ---------- Sellos ----------
sel = wb.create_sheet('Sellos')
cab2 = ['WhatsApp', 'Nombre', 'Sellos', 'Faltan para llavero', 'Faltan para cápsula', 'Llavero canjeado', 'Cápsula canjeada', 'Estado']
for j, h in enumerate(cab2, start=1):
    c = sel.cell(1, j, h); c.font = f(True, 'FFFFFF'); c.fill = CABECERA; c.alignment = Alignment(horizontal='center'); c.border = borde
sel['A1'].comment = Comment('Escribe aquí una vez el WhatsApp de cada cliente (igual que en Pedidos). El resto se calcula.', 'LlaveCorp')
M = 300
for i in range(2, M + 2):
    a = sel.cell(i, 1); a.font = f(color=AZUL); a.fill = AMARILLO; a.border = borde
    sel.cell(i, 2, f'=IF(A{i}="","",IFERROR(INDEX(Pedidos!$C$2:$C${N + 1},MATCH(A{i},Pedidos!$B$2:$B${N + 1},0)),""))')
    sel.cell(i, 3, f'=IF(A{i}="","",SUMIFS(Pedidos!$J$2:$J${N + 1},Pedidos!$B$2:$B${N + 1},A{i}))')
    sel.cell(i, 4, f'=IF(A{i}="","",MAX(0,Config!$B$2-C{i}))')
    sel.cell(i, 5, f'=IF(A{i}="","",MAX(0,Config!$B$3-C{i}))')
    sel.cell(i, 6, f'=IF(A{i}="","",IF(COUNTIFS(Pedidos!$B$2:$B${N + 1},A{i},Pedidos!$D$2:$D${N + 1},"Canje llavero")>0,"Sí","No"))')
    sel.cell(i, 7, f'=IF(A{i}="","",IF(COUNTIFS(Pedidos!$B$2:$B${N + 1},A{i},Pedidos!$D$2:$D${N + 1},"Canje cápsula")>0,"Sí","No"))')
    sel.cell(i, 8, f'=IF(A{i}="","",IF(AND(C{i}>=Config!$B$3,G{i}="No"),"¡Le toca abrir la cápsula!",IF(AND(C{i}>=Config!$B$2,F{i}="No"),"¡Le toca llavero con nombre!",IF(C{i}<Config!$B$2,"Faltan "&D{i}&" para el llavero","Faltan "&E{i}&" para la cápsula"))))')
    for j in range(2, 9):
        sel.cell(i, j).font = f(); sel.cell(i, j).border = borde
sel['A2'] = '528112345678'; sel['A3'] = '528187654321'
for col, w in zip('ABCDEFGH', [15, 18, 8, 18, 18, 16, 16, 32]): sel.column_dimensions[col].width = w
sel.freeze_panes = 'A2'

# ---------- Resumen ----------
res = wb.create_sheet('Resumen')
res['A1'] = 'Resumen'; res['A1'].font = f(True, size=12)
items = [
    ('Pedidos registrados (ventas)', f'=COUNTIF(Pedidos!D2:D{N + 1},"Venta")'),
    ('Piezas vendidas', f'=SUMIFS(Pedidos!F2:F{N + 1},Pedidos!D2:D{N + 1},"Venta")'),
    ('Total vendido ($)', f'=SUMIFS(Pedidos!G2:G{N + 1},Pedidos!D2:D{N + 1},"Venta")'),
    ('Pedidos desde la página', f'=COUNTIFS(Pedidos!D2:D{N + 1},"Venta",Pedidos!I2:I{N + 1},"Página")'),
    ('Llaveros con nombre regalados', f'=COUNTIF(Pedidos!D2:D{N + 1},"Canje llavero")'),
    ('Cápsulas abiertas', f'=COUNTIF(Pedidos!D2:D{N + 1},"Canje cápsula")'),
    ('Clientes a los que ya les toca premio', f'=COUNTIF(Sellos!H2:H{M + 1},"¡Le toca*")'),
]
for i, (etq, formula) in enumerate(items, start=2):
    res.cell(i, 1, etq).font = f(); c = res.cell(i, 2, formula); c.font = f(True)
res['B4'].number_format = '$#,##0'
res.column_dimensions['A'].width = 40; res.column_dimensions['B'].width = 14

wb.move_sheet('Pedidos', offset=-1)
wb.save('marketing/registro-pedidos.xlsx')
print('marketing/registro-pedidos.xlsx generado')
